"""Tabela de prazos de liquidez por ativo (aba `Fundos` da planilha).

A aba tinha duas tabelas lado a lado:

- `A:D` -- Base Origem, Codigo Fundo, CNPJ, Prazo. E a unica lida pelas macros
  (`CalcularDias`, `CalcularDiferencaDias` e `MontarRelatorioLiquidez` casam o
  `Codigo_Ativo` da Carteira contra a coluna B e pegam a coluna D).
- `K:L` -- 14 ativos liquidos com prazo proprio, nao lida por nenhuma macro.

As duas foram unificadas em uma tabela so, `config/prazos_ativos.csv`, com
`A:D` prevalecendo nas divergencias (decisao de projeto). Onde `K:L` divergia,
o valor antigo fica registrado em `observacao` para nao se perder.

Chaves de busca (confirmadas contra a carteira extraida):

- `codigo_ativo` (coluna B) casa com `Carteira.Codigo_Ativo` -- chave principal
- `codigo_op`    (coluna A) casa com `Carteira.Codigo_Op`    -- chave alternativa
"""

from __future__ import annotations

import csv
import logging
from dataclasses import dataclass, field
from pathlib import Path

from .cvm import formatar_cnpj, so_digitos

log = logging.getLogger(__name__)

ARQUIVO_PADRAO = Path("config") / "prazos_ativos.csv"

# Codigo de "sem prazo de resgate" (fundo fechado). A planilha usava 999 e 9999
# para a mesma coisa; foram unificados em 999.
SEM_PRAZO = 999
EQUIVALENTES_SEM_PRAZO = (999, 9999)

COLUNAS = ("codigo_op", "codigo_ativo", "cnpj", "prazo", "origem_prazo", "observacao")

# origem_prazo:
#   cadastro -> o valor veio do pz_resg da ficha do fundo (era VLOOKUP na planilha)
#   fixo     -> o valor foi digitado direto na tabela
#   ""       -> sem prazo definido (os antigos #N/A)
ORIGENS = ("cadastro", "fixo", "")


@dataclass
class Prazo:
    codigo_op: str
    codigo_ativo: str
    cnpj: str | None = None
    prazo: int | None = None
    origem_prazo: str = ""
    observacao: str = ""

    @property
    def sem_prazo(self) -> bool:
        return self.prazo is None

    @property
    def fechado(self) -> bool:
        return self.prazo == SEM_PRAZO

    def to_row(self) -> dict:
        return {
            "codigo_op": self.codigo_op,
            "codigo_ativo": self.codigo_ativo,
            "cnpj": self.cnpj or "",
            "prazo": "" if self.prazo is None else str(self.prazo),
            "origem_prazo": self.origem_prazo,
            "observacao": self.observacao,
        }

    @classmethod
    def from_row(cls, row: dict) -> "Prazo":
        prazo = (row.get("prazo") or "").strip()
        return cls(
            codigo_op=(row.get("codigo_op") or "").strip(),
            codigo_ativo=(row.get("codigo_ativo") or "").strip(),
            cnpj=(row.get("cnpj") or "").strip() or None,
            prazo=int(float(prazo)) if prazo else None,
            origem_prazo=(row.get("origem_prazo") or "").strip(),
            observacao=(row.get("observacao") or "").strip(),
        )


class TabelaPrazos:
    """Lookup de prazo por codigo de ativo, com fallback pelo codigo da operacao."""

    def __init__(self, registros: list[Prazo] | None = None):
        self.registros: list[Prazo] = registros or []
        self._por_ativo: dict[str, Prazo] = {}
        self._por_op: dict[str, Prazo] = {}
        self._indexar()

    def _indexar(self) -> None:
        self._por_ativo, self._por_op = {}, {}
        for registro in self.registros:
            # Primeira ocorrencia vence, igual ao Match do VBA.
            if registro.codigo_ativo:
                self._por_ativo.setdefault(registro.codigo_ativo.strip().upper(), registro)
            if registro.codigo_op:
                self._por_op.setdefault(registro.codigo_op.strip().upper(), registro)

    # ---- persistencia -------------------------------------------------------
    @classmethod
    def carregar(cls, caminho: Path = ARQUIVO_PADRAO) -> "TabelaPrazos":
        caminho = Path(caminho)
        if not caminho.exists():
            raise FileNotFoundError(
                f"{caminho} nao existe. Rode 'python -m liquidez prazos importar'."
            )
        with caminho.open(encoding="utf-8-sig", newline="") as arq:
            registros = [Prazo.from_row(linha) for linha in csv.DictReader(arq, delimiter=";")]
        log.info("Tabela de prazos carregada: %d ativos", len(registros))
        return cls(registros)

    def gravar(self, caminho: Path = ARQUIVO_PADRAO) -> Path:
        caminho = Path(caminho)
        caminho.parent.mkdir(parents=True, exist_ok=True)
        with caminho.open("w", encoding="utf-8-sig", newline="") as arq:
            escritor = csv.DictWriter(arq, fieldnames=list(COLUNAS), delimiter=";")
            escritor.writeheader()
            for registro in sorted(self.registros, key=lambda r: r.codigo_ativo.upper()):
                escritor.writerow(registro.to_row())
        return caminho

    # ---- consulta -----------------------------------------------------------
    def obter(self, codigo: str | None) -> Prazo | None:
        """Busca pelo codigo do ativo; nao achando, tenta pelo codigo da operacao."""
        if not codigo:
            return None
        chave = str(codigo).strip().upper()
        return self._por_ativo.get(chave) or self._por_op.get(chave)

    def prazo_de(self, codigo: str | None) -> int | None:
        registro = self.obter(codigo)
        return registro.prazo if registro else None

    def aplicar(self, df, coluna_ativo: str = "Codigo_Ativo", coluna_op: str | None = None):
        """Devolve uma Series com o prazo de cada linha do DataFrame.

        Por padrao reproduz o Match do VBA, que casa SO pelo `Codigo_Ativo`
        (coluna B da aba). Linhas sem correspondencia ficam nulas -- e o caso
        em que o VBA calcula o prazo pelo DateDiff das datas.

        `coluna_op` habilita o fallback pelo codigo da operacao (coluna A).
        Use com cuidado: a coluna A tem codigos genericos como "A", que casam
        com linhas de CC, CPR e Tesouraria e devolveriam prazo 0 indevidamente.
        """
        import pandas as pd

        if df.empty:
            return pd.Series(dtype="Float64")

        prazos = df[coluna_ativo].map(lambda c: self.prazo_por_ativo(c))
        if coluna_op and coluna_op in df.columns:
            faltando = prazos.isna()
            prazos = prazos.where(
                ~faltando, df[coluna_op].map(lambda c: self.prazo_por_op(c))
            )
        return prazos.astype("Float64")

    def prazo_por_ativo(self, codigo: str | None) -> int | None:
        """Busca estrita pelo codigo do ativo (coluna B), como o VBA faz."""
        if not codigo:
            return None
        registro = self._por_ativo.get(str(codigo).strip().upper())
        return registro.prazo if registro else None

    def prazo_por_op(self, codigo: str | None) -> int | None:
        """Busca pelo codigo da operacao (coluna A)."""
        if not codigo:
            return None
        registro = self._por_op.get(str(codigo).strip().upper())
        return registro.prazo if registro else None

    # ---- validacao ----------------------------------------------------------
    def validar(self, cadastro=None) -> list[str]:
        """Problemas encontrados na tabela (lista vazia = tudo certo)."""
        problemas: list[str] = []
        vistos_ativo: dict[str, str] = {}
        vistos_op: dict[str, str] = {}

        por_cnpj = {}
        if cadastro is not None:
            por_cnpj = {f.cnpj_numerico: f for f in cadastro.fundos.values()}

        for registro in self.registros:
            rotulo = registro.codigo_ativo or registro.codigo_op or "(sem codigo)"

            if not registro.codigo_ativo:
                problemas.append(f"ERRO   {rotulo}: sem codigo_ativo (a chave de busca)")

            chave = registro.codigo_ativo.strip().upper()
            if chave and chave in vistos_ativo:
                problemas.append(
                    f"AVISO  {rotulo}: codigo_ativo duplicado (a busca usa o primeiro, "
                    f"prazo={vistos_ativo[chave]})"
                )
            vistos_ativo[chave] = str(registro.prazo)

            chave_op = registro.codigo_op.strip().upper()
            if chave_op and chave_op in vistos_op:
                problemas.append(f"AVISO  {rotulo}: codigo_op duplicado com {vistos_op[chave_op]}")
            vistos_op[chave_op] = rotulo

            if registro.origem_prazo not in ORIGENS:
                problemas.append(f"ERRO   {rotulo}: origem_prazo invalida ({registro.origem_prazo!r})")

            if registro.prazo is None:
                problemas.append(f"AVISO  {rotulo}: sem prazo definido")
            elif registro.prazo < 0:
                problemas.append(f"ERRO   {rotulo}: prazo negativo ({registro.prazo})")
            elif registro.prazo in EQUIVALENTES_SEM_PRAZO and registro.prazo != SEM_PRAZO:
                problemas.append(f"AVISO  {rotulo}: prazo {registro.prazo} deveria ser {SEM_PRAZO}")

            if registro.cnpj:
                fundo = por_cnpj.get(so_digitos(registro.cnpj))
                if cadastro is not None and fundo is None:
                    problemas.append(f"AVISO  {rotulo}: CNPJ {registro.cnpj} nao esta no cadastro")
                elif fundo is not None and registro.prazo is not None and fundo.pz_resg is not None:
                    if registro.prazo != fundo.pz_resg:
                        problemas.append(
                            f"AVISO  {rotulo}: prazo {registro.prazo} difere do cadastro "
                            f"({fundo.nome_interno}: pz_resg={fundo.pz_resg})"
                        )
        return problemas

    def divergentes_do_cadastro(self, cadastro) -> list[tuple[Prazo, int]]:
        """Registros com origem 'cadastro' cujo prazo saiu de sincronia com a ficha."""
        por_cnpj = {f.cnpj_numerico: f for f in cadastro.fundos.values()}
        fora: list[tuple[Prazo, int]] = []
        for registro in self.registros:
            if registro.origem_prazo != "cadastro" or not registro.cnpj:
                continue
            fundo = por_cnpj.get(so_digitos(registro.cnpj))
            if fundo and fundo.pz_resg is not None and fundo.pz_resg != registro.prazo:
                fora.append((registro, fundo.pz_resg))
        return fora


# ---------------------------------------------------------------------------
# Importacao da planilha
# ---------------------------------------------------------------------------

def _inteiro(valor) -> int | None:
    if valor is None or valor == "":
        return None
    try:
        numero = int(float(valor))
    except (TypeError, ValueError):
        return None
    return SEM_PRAZO if numero in EQUIVALENTES_SEM_PRAZO else numero


def importar_da_planilha(caminho_xlsm: Path, aba: str = "Fundos") -> tuple[TabelaPrazos, dict]:
    """Le a aba `Fundos` e devolve (tabela unificada, relatorio da importacao)."""
    from openpyxl import load_workbook

    # data_only=True traz o valor calculado; a segunda leitura diz se a celula
    # do prazo era formula (VLOOKUP no Config_Fundos) ou valor digitado.
    wb_valores = load_workbook(caminho_xlsm, read_only=True, data_only=True, keep_vba=False)
    wb_formulas = load_workbook(caminho_xlsm, read_only=True, data_only=False, keep_vba=False)
    try:
        ws_valores, ws_formulas = wb_valores[aba], wb_formulas[aba]
        linhas_valores = list(ws_valores.iter_rows(min_row=2, max_col=12, values_only=True))
        linhas_formulas = list(ws_formulas.iter_rows(min_row=2, max_col=12, values_only=True))
    finally:
        wb_valores.close()
        wb_formulas.close()

    # Tabela K:L (colunas 11 e 12), usada so para registrar as divergencias.
    tabela_kl: dict[str, int | None] = {}
    for linha in linhas_valores:
        codigo = linha[10] if len(linha) > 10 else None
        if codigo:
            tabela_kl[str(codigo).strip().upper()] = _inteiro(linha[11] if len(linha) > 11 else None)

    registros: list[Prazo] = []
    relatorio = {"lidos": 0, "sem_prazo": 0, "unificados_9999": 0, "divergencias_kl": [], "duplicados": []}
    vistos: set[str] = set()

    for indice, linha in enumerate(linhas_valores):
        codigo_op = str(linha[0]).strip() if linha[0] is not None else ""
        codigo_ativo = str(linha[1]).strip() if linha[1] is not None else ""
        if not codigo_op and not codigo_ativo:
            continue
        relatorio["lidos"] += 1

        cnpj_bruto = str(linha[2]).strip() if linha[2] is not None else ""
        cnpj = formatar_cnpj(cnpj_bruto) if cnpj_bruto else None

        bruto = linha[3]
        if isinstance(bruto, str) and "N/A" in bruto.upper():
            bruto = None
        prazo = _inteiro(bruto)
        try:
            if bruto is not None and int(float(bruto)) == 9999:
                relatorio["unificados_9999"] += 1
        except (TypeError, ValueError):
            pass

        celula_formula = linhas_formulas[indice][3] if indice < len(linhas_formulas) else None
        era_formula = isinstance(celula_formula, str) and celula_formula.startswith("=")
        origem = "" if prazo is None else ("cadastro" if era_formula else "fixo")
        if prazo is None:
            relatorio["sem_prazo"] += 1

        observacoes = []
        chave = codigo_ativo.upper()
        prazo_kl = tabela_kl.get(chave)
        if prazo_kl is not None and prazo_kl != prazo:
            observacoes.append(f"tabela K:L da planilha indicava prazo {prazo_kl}")
            relatorio["divergencias_kl"].append((codigo_ativo, prazo_kl, prazo))
        if chave and chave in vistos:
            observacoes.append("codigo_ativo repetido na planilha; a busca usa o primeiro")
            relatorio["duplicados"].append(codigo_ativo)
        vistos.add(chave)

        registros.append(
            Prazo(
                codigo_op=codigo_op,
                codigo_ativo=codigo_ativo,
                cnpj=cnpj,
                prazo=prazo,
                origem_prazo=origem,
                observacao="; ".join(observacoes),
            )
        )

    return TabelaPrazos(registros), relatorio
