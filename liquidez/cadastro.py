"""Cadastro unico de fundos: uma ficha JSON por fundo em `config/fundos/`.

Cada ficha tem quatro blocos e a separacao entre eles e o que faz o cadastro
funcionar:

- `cvm`        sincronizado do portal, SOBRESCRITO a cada sync, nao editar a mao
- `liquidez`   pz_resg / cc, mantidos por nos, o sync NUNCA toca
- `integracao` codinome (Beyond) e nome_produto (MOV)
- `controle`   ativo/inativo, observacoes, carimbo de atualizacao
"""

from __future__ import annotations

import datetime as dt
import json
import logging
from dataclasses import dataclass, field
from pathlib import Path

from .cvm import SITUACAO_CANCELADO, formatar_cnpj, so_digitos

log = logging.getLogger(__name__)

DIR_CADASTRO = Path("config") / "fundos"
DIR_REMOVIDOS = DIR_CADASTRO / "_removidos"

# Situacoes de registro que ainda nao vigoram, ou nao vigoram mais. Um pedido
# `Em Analise` e intencao protocolada, nao fato -- descrever o fundo por ele
# seria antecipar deferimento da CVM. Comparadas em maiusculas, sem acento
# apenas onde a CVM ja escreve assim.
SITUACOES_NAO_VIGENTES = frozenset({
    "EM ANÁLISE", "EM ANALISE", "CANCELADA", "CANCELADO",
})


def validar_cnpj(cnpj: str | None) -> bool:
    """Valida os dois digitos verificadores do CNPJ."""
    d = so_digitos(cnpj)
    if len(d) != 14 or len(set(d)) == 1:
        return False
    for tamanho in (12, 13):
        pesos = list(range(tamanho - 7, 1, -1)) + list(range(9, 1, -1))
        soma = sum(int(d[i]) * pesos[i] for i in range(tamanho))
        resto = soma % 11
        digito = 0 if resto < 2 else 11 - resto
        if int(d[tamanho]) != digito:
            return False
    return True


@dataclass
class Fundo:
    """Ficha de um fundo. Espelha o JSON gravado em disco."""

    cnpj: str
    nome_interno: str
    cvm: dict = field(default_factory=dict)
    liquidez: dict = field(default_factory=dict)
    integracao: dict = field(default_factory=dict)
    controle: dict = field(default_factory=dict)

    # ---- atalhos de leitura -------------------------------------------------
    @property
    def cnpj_numerico(self) -> str:
        return so_digitos(self.cnpj)

    @property
    def administrador(self) -> str | None:
        """Administrador vigente. A CVM manda (decisao de projeto)."""
        adm = (self.cvm or {}).get("administrador") or {}
        return adm.get("nome")

    @property
    def gestor(self) -> str | None:
        """Gestor vigente, pela CVM -- a mesma fonte do administrador.

        O `riscogc` tambem tem gestor, em `HOD_TB_CADASTRO_FUNDOS`, mas as duas
        fontes discordam: em 31/07/2026 os dois ECOAGRO apareciam TROCADOS entre
        si e o SALUS trazia uma administradora no campo de gestor. A CVM e o
        registro oficial e ja esta na ficha, entao e ela que vale aqui.
        """
        g = (self.cvm or {}).get("gestor") or {}
        return g.get("nome")

    @property
    def cnpj_gestor(self) -> str | None:
        return ((self.cvm or {}).get("gestor") or {}).get("cnpj")

    @property
    def tipo_fundo(self) -> str | None:
        return (self.cvm or {}).get("tipo_fundo")

    @property
    def situacao_cvm(self) -> str | None:
        return (self.cvm or {}).get("situacao")

    @property
    def forma_condominio(self) -> str | None:
        """Aberto ou Fechado, pela CVM. None quando a ficha nao diz.

        A CVM registra a forma em dois niveis e nem sempre nos dois. Na carga de
        31/07/2026, 84 dos 100 fundos trazem no nivel da CLASSE e os 16 restantes
        so nas SUBCLASSES -- por isso a busca desce, em vez de devolver vazio
        para um sexto do cadastro.

        A classe vence a subclasse quando as duas existem: e o nivel do fundo, e
        a subclasse e recorte dentro dele.

        Registro que ainda nao vigora nao define a forma do fundo. O MURALE FIDC
        tem DUAS subclasses "SENIOR": uma Fechada `Em Analise` e outra Aberta
        `Fase Pre-Operacional` -- e uma re-protocolizacao, e a versao antiga
        segue em analise na CVM. Contando as duas, o fundo aparecia como
        "Aberto / Fechado" por causa de um pedido que ainda nao foi deferido.
        `SITUACOES_NAO_VIGENTES` corta esses; em 31/07/2026 sao 2 subclasses no
        cadastro inteiro, e o unico fundo afetado e o MURALE.

        Se o corte levar tudo, o valor bruto volta: e melhor mostrar a forma de
        um registro em analise que nao mostrar forma nenhuma.

        Divergencia que sobreviva ao corte nao e erro nem escolha nossa -- vira
        "Aberto / Fechado", porque eleger um vencedor esconderia justamente a
        informacao que importa: a de o fundo nao ter forma unica.
        """
        def formas(itens: list, filtrar: bool) -> list[str]:
            saida = []
            for it in itens:
                valor = it.get("forma_condominio")
                if not valor:
                    continue
                situacao = str(it.get("situacao") or "").strip().upper()
                if filtrar and situacao in SITUACOES_NAO_VIGENTES:
                    continue
                saida.append(str(valor).strip())
            return saida

        classes = ((self.cvm or {}).get("classes")) or []
        for nivel in ("classe", "subclasse"):
            if nivel == "classe":
                itens = list(classes)
            else:
                itens = [s for c in classes for s in (c.get("subclasses") or [])]
            unicos = sorted(set(formas(itens, filtrar=True)))
            if not unicos:
                unicos = sorted(set(formas(itens, filtrar=False)))
            if unicos:
                return " / ".join(unicos)
        return None

    @property
    def data_inicio_situacao(self) -> str | None:
        """Desde quando o fundo esta na situacao atual, em ISO.

        Ganha sentido quando a situacao nao e normal: no MANI e no KAPAMA ela e
        a data em que entraram em liquidacao -- e bate exatamente com a ultima
        data em que cada um aparece na carteira.
        """
        return (self.cvm or {}).get("data_inicio_situacao") or None

    @property
    def em_funcionamento(self) -> bool:
        """A situacao da CVM e a normal de um fundo operando."""
        return str(self.situacao_cvm or "").strip().upper() == "EM FUNCIONAMENTO NORMAL"

    @property
    def data_constituicao(self) -> str | None:
        """Quando o fundo foi constituido, em ISO. Preenchida nos 100 fundos.

        E a mais antiga das datas da ficha e a que mais se aproxima de "quando o
        fundo nasceu". Nao confundir com `classe.data_inicio`: aquela e o
        registro da CLASSE, criada pela RCVM 175, e por isso nenhuma e anterior a
        31/10/2023 -- o ZR11, constituido em 2013, tem classe iniciada em 2025.
        Usar a da classe como idade do fundo erraria por uma decada.

        Tambem nao e inicio de OPERACAO: a CVM nao registra isso. O mais proximo
        disso na base e a primeira data da serie de cotas.
        """
        return (self.cvm or {}).get("data_constituicao") or None

    @property
    def pz_resg(self) -> int | None:
        return (self.liquidez or {}).get("pz_resg")

    @property
    def cc(self) -> int | None:
        return (self.liquidez or {}).get("cc")

    @property
    def codinome_beyond(self) -> str | None:
        return (self.integracao or {}).get("codinome_beyond")

    @property
    def nome_produto_mov(self) -> str:
        return (self.integracao or {}).get("nome_produto_mov") or self.nome_interno

    @property
    def ativo(self) -> bool:
        return bool((self.controle or {}).get("ativo", True))

    @property
    def incluir_no_lote(self) -> bool:
        return bool((self.controle or {}).get("incluir_no_lote", True))

    # ---- serializacao -------------------------------------------------------
    def to_dict(self) -> dict:
        return {
            "cnpj": self.cnpj,
            "cnpj_numerico": self.cnpj_numerico,
            "nome_interno": self.nome_interno,
            "cvm": self.cvm,
            "liquidez": self.liquidez,
            "integracao": self.integracao,
            "controle": self.controle,
        }

    @classmethod
    def from_dict(cls, dados: dict) -> "Fundo":
        return cls(
            cnpj=dados["cnpj"],
            nome_interno=dados["nome_interno"],
            cvm=dados.get("cvm") or {},
            liquidez=dados.get("liquidez") or {},
            integracao=dados.get("integracao") or {},
            controle=dados.get("controle") or {},
        )

    def marcar_atualizado(self) -> None:
        self.controle["atualizado_em"] = dt.date.today().isoformat()


class Cadastro:
    """Colecao de fichas em disco."""

    def __init__(self, diretorio: Path = DIR_CADASTRO):
        self.diretorio = Path(diretorio)
        self.fundos: dict[str, Fundo] = {}

    # ---- persistencia -------------------------------------------------------
    def carregar(self) -> "Cadastro":
        self.fundos = {}
        if not self.diretorio.exists():
            return self
        for caminho in sorted(self.diretorio.glob("*.json")):
            dados = json.loads(caminho.read_text(encoding="utf-8"))
            fundo = Fundo.from_dict(dados)
            self.fundos[fundo.cnpj_numerico] = fundo
        log.info("Cadastro carregado: %d fundos", len(self.fundos))
        return self

    def caminho_de(self, fundo: Fundo) -> Path:
        return self.diretorio / f"{fundo.cnpj_numerico}.json"

    def gravar(self, fundo: Fundo) -> Path:
        self.diretorio.mkdir(parents=True, exist_ok=True)
        caminho = self.caminho_de(fundo)
        caminho.write_text(
            json.dumps(fundo.to_dict(), ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        self.fundos[fundo.cnpj_numerico] = fundo
        return caminho

    def gravar_todos(self) -> None:
        for fundo in self.fundos.values():
            self.gravar(fundo)

    def remover(self, cnpj: str, motivo: str) -> Path | None:
        """Move a ficha para `_removidos/`, preservando o rastro."""
        chave = so_digitos(cnpj)
        fundo = self.fundos.pop(chave, None)
        if fundo is None:
            return None
        DIR_REMOVIDOS.mkdir(parents=True, exist_ok=True)
        dados = fundo.to_dict()
        dados["removido"] = {"motivo": motivo, "data": dt.date.today().isoformat()}
        destino = DIR_REMOVIDOS / f"{chave}.json"
        destino.write_text(json.dumps(dados, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        origem = self.diretorio / f"{chave}.json"
        if origem.exists():
            origem.unlink()
        return destino

    # ---- consulta -----------------------------------------------------------
    def obter(self, identificador: str) -> Fundo | None:
        """Localiza por CNPJ (com ou sem mascara) ou por nome interno."""
        chave = so_digitos(identificador)
        if chave and chave in self.fundos:
            return self.fundos[chave]
        alvo = identificador.strip().upper()
        for fundo in self.fundos.values():
            if fundo.nome_interno.strip().upper() == alvo:
                return fundo
        return None

    def filtrar(
        self,
        adm: str | None = None,
        tipo: str | None = None,
        situacao: str | None = None,
        incluir_inativos: bool = False,
    ) -> list[Fundo]:
        selecionados = []
        for fundo in sorted(self.fundos.values(), key=lambda f: f.nome_interno):
            if not incluir_inativos and not (fundo.ativo and fundo.incluir_no_lote):
                continue
            if adm and adm.upper() not in (fundo.administrador or "").upper():
                continue
            if tipo and tipo.upper() != (fundo.tipo_fundo or "").upper():
                continue
            if situacao and situacao.upper() not in (fundo.situacao_cvm or "").upper():
                continue
            selecionados.append(fundo)
        return selecionados

    # ---- validacao ----------------------------------------------------------
    def validar(self) -> list[str]:
        """Devolve a lista de problemas encontrados (vazia = cadastro ok)."""
        from .regras_adm import obter as obter_regra

        problemas: list[str] = []
        vistos_cnpj: dict[str, str] = {}
        vistos_nome: dict[str, str] = {}

        for fundo in self.fundos.values():
            rotulo = f"{fundo.nome_interno} ({fundo.cnpj})"

            if not validar_cnpj(fundo.cnpj):
                problemas.append(f"ERRO   {rotulo}: CNPJ invalido (digito verificador)")
            if fundo.cnpj_numerico in vistos_cnpj:
                problemas.append(f"ERRO   {rotulo}: CNPJ duplicado com {vistos_cnpj[fundo.cnpj_numerico]}")
            vistos_cnpj[fundo.cnpj_numerico] = rotulo

            nome = fundo.nome_interno.strip().upper()
            if nome in vistos_nome:
                problemas.append(f"ERRO   {rotulo}: nome duplicado com {vistos_nome[nome]}")
            vistos_nome[nome] = rotulo

            if fundo.pz_resg is None:
                problemas.append(f"ERRO   {rotulo}: liquidez.pz_resg nao preenchido")
            if fundo.cc is None:
                problemas.append(f"AVISO  {rotulo}: liquidez.cc nao preenchido")
            if not fundo.cvm:
                problemas.append(f"AVISO  {rotulo}: sem dados da CVM (rodar sync-cvm)")
            if fundo.situacao_cvm == SITUACAO_CANCELADO:
                problemas.append(f"AVISO  {rotulo}: cancelado na CVM, deveria sair do cadastro")
            if not fundo.codinome_beyond:
                problemas.append(f"AVISO  {rotulo}: sem codinome_beyond, a base Beyond sera pulada")

            regra = obter_regra(fundo.administrador)
            if not regra.implementado:
                problemas.append(
                    f"AVISO  {rotulo}: administrador sem tratamento de liquidez ({fundo.administrador})"
                )
        return problemas


# ---------------------------------------------------------------------------
# Carga inicial a partir da planilha
# ---------------------------------------------------------------------------

def ler_config_fundos(caminho_xlsm: Path) -> list[dict]:
    """Le a aba Config_Fundos: A=NomeFundo, B=DocFundo, C=PzResg, D=CC, E=Administrador."""
    from openpyxl import load_workbook

    wb = load_workbook(caminho_xlsm, read_only=True, data_only=True, keep_vba=False)
    try:
        ws = wb["Config_Fundos"]
        registros = []
        for linha in ws.iter_rows(min_row=2, max_col=5, values_only=True):
            nome = (linha[0] or "").strip() if isinstance(linha[0], str) else linha[0]
            if not nome:
                continue
            registros.append(
                {
                    "nome": str(nome).strip(),
                    "cnpj": str(linha[1] or "").strip(),
                    "pz_resg": linha[2],
                    "cc": linha[3],
                    "administrador_planilha": str(linha[4] or "").strip() or None,
                }
            )
        return registros
    finally:
        wb.close()


def ler_codinomes(caminho_xlsm: Path) -> dict[str, str]:
    """Le os codinomes do Beyond da aba Painel (S=nome, T=CNPJ, U=codinome)."""
    from openpyxl import load_workbook

    wb = load_workbook(caminho_xlsm, read_only=True, data_only=True, keep_vba=False)
    try:
        ws = wb["Painel"]
        codinomes: dict[str, str] = {}
        for linha in ws.iter_rows(min_row=2, min_col=19, max_col=21, values_only=True):
            _, cnpj, codinome = linha
            if cnpj and codinome:
                codinomes[so_digitos(str(cnpj))] = str(codinome).strip()
        return codinomes
    finally:
        wb.close()


def _inteiro(valor) -> int | None:
    if valor is None or valor == "":
        return None
    try:
        return int(float(valor))
    except (TypeError, ValueError):
        return None


def montar_ficha(registro: dict, bloco_cvm: dict | None, codinome: str | None) -> Fundo:
    """Cria a ficha de um fundo a partir da planilha + CVM."""
    hoje = dt.date.today().isoformat()
    return Fundo(
        cnpj=formatar_cnpj(registro["cnpj"]) or registro["cnpj"],
        nome_interno=registro["nome"],
        cvm=bloco_cvm or {},
        liquidez={"pz_resg": _inteiro(registro.get("pz_resg")), "cc": _inteiro(registro.get("cc"))},
        integracao={
            "codinome_beyond": codinome,
            "nome_produto_mov": registro["nome"],
        },
        controle={
            "ativo": True,
            "incluir_no_lote": True,
            "observacoes": "",
            "administrador_planilha_origem": registro.get("administrador_planilha"),
            "atualizado_em": hoje,
        },
    )
