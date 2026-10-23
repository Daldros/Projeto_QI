"""Gera o Relatorio de Liquidez formatado em Excel, um arquivo por fundo.

Reproduz o layout da aba `Relatorio_Liquidez`: parametros, tabela por ativo,
indicadores das pontas ativa e passiva e o enquadramento colorido.

Consome as saidas ja validadas (etapa 9 + indicadores); nao recalcula nada.
"""

from __future__ import annotations

import datetime as dt
import logging
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .indicadores import Indicadores

log = logging.getLogger(__name__)

# --- Paleta -----------------------------------------------------------------
AZUL = "1F4E79"
AZUL_CLARO = "D6E4F0"
CINZA = "F2F2F2"
CINZA_BORDA = "BFBFBF"
VERDE_FUNDO = "C6EFCE"
VERDE_TEXTO = "006100"
VERMELHO_FUNDO = "FFC7CE"
VERMELHO_TEXTO = "9C0006"
NEUTRO_FUNDO = "FFEB9C"
NEUTRO_TEXTO = "9C6500"

MOEDA = '#,##0.00'
PCT = '0.00%'

_fino = Side(style="thin", color=CINZA_BORDA)
BORDA = Border(left=_fino, right=_fino, top=_fino, bottom=_fino)


def _fmt_data(valor) -> str:
    if isinstance(valor, dt.date):
        return valor.strftime("%d/%m/%Y")
    try:
        return dt.date.fromisoformat(str(valor)).strftime("%d/%m/%Y")
    except (ValueError, TypeError):
        return str(valor)


class _Escritor:
    """Auxiliar para escrever celulas com estilo, controlando a linha atual."""

    def __init__(self, ws):
        self.ws = ws
        self.linha = 1

    def pula(self, n=1):
        self.linha += n

    def titulo(self, texto, sub=None):
        c = self.ws.cell(self.linha, 1, texto)
        c.font = Font(bold=True, size=16, color=AZUL)
        self.linha += 1
        if sub:
            c = self.ws.cell(self.linha, 1, sub)
            c.font = Font(italic=True, size=10, color="808080")
            self.linha += 1

    def secao(self, texto, largura=5):
        self.linha += 1
        for col in range(1, largura + 1):
            c = self.ws.cell(self.linha, col)
            c.fill = PatternFill("solid", fgColor=AZUL)
        c = self.ws.cell(self.linha, 1, texto)
        c.font = Font(bold=True, color="FFFFFF", size=11)
        self.ws.merge_cells(start_row=self.linha, start_column=1,
                            end_row=self.linha, end_column=largura)
        self.linha += 1

    def par(self, rotulo, valor, fmt=None, col=1):
        r = self.ws.cell(self.linha, col, rotulo)
        r.font = Font(bold=True)
        r.fill = PatternFill("solid", fgColor=CINZA)
        v = self.ws.cell(self.linha, col + 1, valor)
        if fmt:
            v.number_format = fmt
        return v

    def enquadramento(self, rotulo, situacao: bool | None, col=1):
        r = self.ws.cell(self.linha, col, rotulo)
        r.font = Font(bold=True)
        r.fill = PatternFill("solid", fgColor=CINZA)
        c = self.ws.cell(self.linha, col + 1)
        if situacao is None:
            c.value, fundo, texto = "n/d", NEUTRO_FUNDO, NEUTRO_TEXTO
        elif situacao:
            c.value, fundo, texto = "Enquadrado", VERDE_FUNDO, VERDE_TEXTO
        else:
            c.value, fundo, texto = "Desenquadrado", VERMELHO_FUNDO, VERMELHO_TEXTO
        c.fill = PatternFill("solid", fgColor=fundo)
        c.font = Font(bold=True, color=texto)
        c.alignment = Alignment(horizontal="center")


def gerar_excel(
    fundo,
    data_posicao,
    relatorio: pd.DataFrame,
    pl_mil: float,
    indicadores: Indicadores,
    destino: Path,
) -> Path:
    """Monta o relatorio de um fundo e grava o .xlsx. Devolve o caminho."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Relatorio_Liquidez"
    ws.sheet_view.showGridLines = False
    e = _Escritor(ws)

    def pct(v):
        return "n/d" if v is None else v

    # --- Cabecalho / parametros -----------------------------------------
    e.titulo("Relatório Risco de Liquidez",
             "Análise de Adequação de Liquidez — Ponta Ativa e Passiva")
    e.secao("Parâmetros")
    e.par("Data", _fmt_data(data_posicao)); e.par("Administrador", fundo.administrador or "-", col=4); e.pula()
    e.par("Fundo", fundo.nome_interno); e.par("CNPJ", fundo.cnpj, col=4); e.pula()
    e.par("PL (R$ mil)", pl_mil, MOEDA)
    e.par("Prazo de cotização (dias)", indicadores.prazo_cotizacao, col=4); e.pula()

    # --- Resultado por Ativo --------------------------------------------
    e.secao("Resultado por Ativo")
    cabec = ["Código Ativo", "Exposição", "Liq. Duration", "Valor Líquido"]
    for j, texto in enumerate(cabec, start=1):
        c = ws.cell(e.linha, j, texto)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL_CLARO)
        c.font = Font(bold=True, color=AZUL)
        c.border = BORDA
    e.linha += 1
    for _, linha in relatorio.iterrows():
        ws.cell(e.linha, 1, linha["codigo_ativo"]).border = BORDA
        ws.cell(e.linha, 2, float(linha["exposicao"])).number_format = MOEDA
        ws.cell(e.linha, 3, linha["liq_duration"]).alignment = Alignment(horizontal="center")
        ws.cell(e.linha, 4, float(linha["valor_liquido"])).number_format = MOEDA
        for j in range(1, 5):
            ws.cell(e.linha, j).border = BORDA
        e.linha += 1
    # Total dos ativos listados -- nao e o PL: a etapa 9 ignora as cotas
    # marcadas como vencidas, que entram no PL mas nao nesta tabela.
    ws.cell(e.linha, 1, "Total dos ativos listados").font = Font(bold=True)
    tc = ws.cell(e.linha, 2, float(pd.to_numeric(relatorio["exposicao"], errors="coerce").sum()))
    tc.number_format = MOEDA; tc.font = Font(bold=True)
    tv = ws.cell(e.linha, 4, float(pd.to_numeric(relatorio["valor_liquido"], errors="coerce").sum()))
    tv.number_format = MOEDA; tv.font = Font(bold=True)
    for j in range(1, 5):
        ws.cell(e.linha, j).fill = PatternFill("solid", fgColor=CINZA)
        ws.cell(e.linha, j).border = BORDA
    e.linha += 1

    ind = indicadores

    # --- Ponta Ativa: Caixa Potencial -----------------------------------
    e.secao("Ponta Ativa — Caixa Potencial")
    e.par("Caixa Potencial (R$ mil)", ind.caixa_potencial_mil, MOEDA); e.pula()
    v = e.par("Caixa / PL", pct(ind.caixa_pl), PCT if ind.caixa_pl is not None else None); e.pula()

    # ANBIMA
    e.par("ANBIMA — Percentil 95% por horizonte (dias)", "")
    e.linha += 1
    ws.cell(e.linha, 1, "Horizonte").font = Font(bold=True)
    ws.cell(e.linha, 1).fill = PatternFill("solid", fgColor=CINZA)
    for k, (h, _, _) in enumerate(ind.anbima, start=2):
        c = ws.cell(e.linha, k, h); c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center")
        c.fill = PatternFill("solid", fgColor=CINZA)
    e.linha += 1
    ws.cell(e.linha, 1, "Percentil 95%").font = Font(bold=True)
    for k, (_, p, _) in enumerate(ind.anbima, start=2):
        c = ws.cell(e.linha, k, p); c.number_format = PCT
        c.alignment = Alignment(horizontal="center")
    e.linha += 1
    ws.cell(e.linha, 1, "Situação").font = Font(bold=True)
    for k, (_, _, enq) in enumerate(ind.anbima, start=2):
        c = ws.cell(e.linha, k, "Enq." if enq else "Desenq.")
        c.fill = PatternFill("solid", fgColor=VERDE_FUNDO if enq else VERMELHO_FUNDO)
        c.font = Font(bold=True, color=VERDE_TEXTO if enq else VERMELHO_TEXTO, size=9)
        c.alignment = Alignment(horizontal="center")
    e.linha += 1

    # --- Ponta Passiva --------------------------------------------------
    e.secao("Ponta Passiva — Concentração de Cotistas")
    e.par("Nº de cotistas", ind.n_cotistas)
    e.par("Maior cotista (R$ mil)", ind.maior_cotista_mil, MOEDA, col=4); e.pula()
    e.par("Cx / Cotista", pct(ind.cx_cotista), PCT if ind.cx_cotista is not None else None)
    e.par("3 maiores (R$ mil)", ind.tres_maiores_cotistas_mil, MOEDA, col=4); e.pula()
    e.enquadramento("Situação (Cx/Cot > 15%)", ind.enq_cotista); e.pula()

    e.secao("Ponta Passiva — Grau de Dispersão do Passivo (GDP)")
    e.par("Média (R$ mil)", ind.media_mil, MOEDA)
    e.par("Desvio (R$ mil)", ind.desvio_mil, MOEDA, col=4); e.pula()
    e.par("GDP (R$ mil)", ind.gdp_mil, MOEDA)
    e.par("Cx / GDP", pct(ind.cx_gdp), PCT if ind.cx_gdp is not None else None, col=4); e.pula()
    e.enquadramento("Situação (Cx/GDP > 15%)", ind.enq_gdp); e.pula()

    e.secao("Ponta Passiva — Distribuidores")
    e.par("Maior distribuidor (R$ mil)", ind.maior_distribuidor_mil, MOEDA)
    e.par("Cx / Distribuidor", pct(ind.cx_distribuidor), PCT if ind.cx_distribuidor is not None else None, col=4)
    e.pula()
    e.enquadramento("Situação (Cx/Dist > 1)", ind.enq_distribuidor); e.pula()

    e.secao("Resgate em Condições Ordinárias (RCO)")
    e.par("Maior resgate diário (R$ mil)", ind.resgate_max_mil, MOEDA)
    e.par("Caixa / Resgate", pct(ind.caixa_resgate), PCT if ind.caixa_resgate is not None else None, col=4)
    e.pula()
    e.enquadramento("Situação (Caixa/Resg > 1)", ind.enq_rco); e.pula()

    # --- Larguras --------------------------------------------------------
    larguras = {1: 30, 2: 18, 3: 14, 4: 22, 5: 16}
    for col, w in larguras.items():
        ws.column_dimensions[get_column_letter(col)].width = w

    ws.cell(e.linha + 1, 1, f"Gerado por python -m liquidez em {dt.date.today().strftime('%d/%m/%Y')}").font = \
        Font(italic=True, size=8, color="808080")

    destino = Path(destino)
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    log.info("Relatorio Excel gerado: %s", destino)
    return destino
